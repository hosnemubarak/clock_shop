import os
from io import BytesIO
from django.http import HttpResponse
from django.template.loader import get_template
from django.conf import settings
from django.utils import timezone
from xhtml2pdf import pisa
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def generate_pdf(template_src, context_dict, filename):
    """
    Renders a Django template into a PDF and returns an HttpResponse.
    """
    template = get_template(template_src)
    
    # Add common context variables for the PDF
    context_dict['generation_date'] = timezone.localtime().strftime('%Y-%m-%d %H:%M:%S')
    context_dict['company_name'] = getattr(settings, 'COMPANY_NAME', 'Clock Shop')
    
    html = template.render(context_dict)
    result = BytesIO()
    
    # Create PDF
    pdf = pisa.CreatePDF(BytesIO(html.encode("UTF-8")), dest=result)
    
    if not pdf.err:
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
        return response
    return HttpResponse('Errors during PDF generation.', status=400)


def generate_excel(filename, title, filters_dict, headers, data, totals=None):
    """
    Generates an Excel file with the given data and returns an HttpResponse.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    
    # Styles
    title_font = Font(size=14, bold=True)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    bold_font = Font(bold=True)
    center_aligned = Alignment(horizontal="center", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    current_row = 1
    
    # Title
    ws.cell(row=current_row, column=1, value=title).font = title_font
    current_row += 2
    
    # Filters
    if filters_dict:
        ws.cell(row=current_row, column=1, value="Applied Filters:").font = bold_font
        current_row += 1
        for key, value in filters_dict.items():
            if value:
                ws.cell(row=current_row, column=1, value=f"{key}:")
                ws.cell(row=current_row, column=2, value=str(value))
                current_row += 1
        current_row += 1
        
    # Generation Date
    ws.cell(row=current_row, column=1, value="Generated On:").font = bold_font
    ws.cell(row=current_row, column=2, value=timezone.localtime().strftime('%Y-%m-%d %H:%M:%S'))
    current_row += 2
    
    # Table Headers
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=current_row, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_aligned
        cell.border = thin_border
    
    current_row += 1
    
    # Data Rows
    for row_data in data:
        for col_num, cell_value in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_num, value=cell_value)
            cell.border = thin_border
        current_row += 1
        
    # Totals Row
    if totals:
        for col_num, total_val in enumerate(totals, 1):
            if total_val is not None and total_val != '':
                cell = ws.cell(row=current_row, column=col_num, value=total_val)
                cell.font = bold_font
                cell.border = thin_border
        current_row += 1
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2)
        if adjusted_width > 50:
            adjusted_width = 50
        ws.column_dimensions[column].width = adjusted_width

    # Return response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response
